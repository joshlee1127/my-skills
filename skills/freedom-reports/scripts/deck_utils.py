"""
Shared helpers for building decks from the Freedom Systems (自由系統) template.

Always start from a COPY of assets/自由系統簡報模板_2021.pptx — never edit the
original in place, and never edit the copy inside ~/.claude/skills/freedom-report/
itself. Copy it into the working/output directory first, then operate on that copy.

Usage is via import from other scripts in this folder, e.g.:
    from deck_utils import list_layouts, delete_slide, duplicate_slide, copy_shape

python-pptx does not support cross-presentation layout copying, slide duplication,
or slide deletion natively — that's why these exist. Everything here only ever
edits one open Presentation object; it never merges two files.
"""
import copy
import sys
from pptx import Presentation
from pptx.util import Emu

# list_layouts()/list_slides() print Traditional Chinese layout/shape names and
# the "© Freedom Systems Inc." footer text. On a Windows console whose active
# code page is cp950 (common on zh-TW machines), plain print() raises
# UnicodeEncodeError on that content instead of printing it — reconfigure to
# UTF-8 so these functions work regardless of how the caller invoked Python.
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        try:
            _stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


def list_layouts(prs_or_path):
    """Print every slide-master layout with its name, index, and placeholder
    idx/type so you can pick the right layout name for add_content_slide()."""
    prs = prs_or_path if hasattr(prs_or_path, "slides") else Presentation(prs_or_path)
    for m_i, master in enumerate(prs.slide_masters):
        print(f"MASTER {m_i}")
        for l_i, layout in enumerate(master.slide_layouts):
            print(f"  [{l_i}] {layout.name}")
            for ph in layout.placeholders:
                print(f"      idx={ph.placeholder_format.idx} type={ph.placeholder_format.type} name={ph.name}")


def list_slides(prs_or_path):
    """Print every existing slide with its layout name and shape names —
    use this to find the slide index of the icon library / example diagrams
    before copying shapes out of them, or before deleting a slide."""
    prs = prs_or_path if hasattr(prs_or_path, "slides") else Presentation(prs_or_path)
    for i, slide in enumerate(prs.slides):
        print(f"slide {i}: layout={slide.slide_layout.name!r}")
        for shape in slide.shapes:
            txt = shape.text_frame.text[:40].replace("\n", " | ") if shape.has_text_frame else ""
            print(f"    id={shape.shape_id} name={shape.name!r} text={txt!r}")


def find_layout(prs, name):
    """Find a slide layout by exact name across all masters. Raises if not found
    — layout names are in Traditional Chinese in this template (e.g. '標題及內容',
    '標題投影片', '章節標題'), run list_layouts() first if unsure."""
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.name == name:
                return layout
    raise KeyError(f"layout {name!r} not found — run list_layouts() to see available names")


def set_placeholder_text(slide, idx, text):
    """Set text on the placeholder with the given idx (from list_layouts output).
    Silently no-ops if the slide has no placeholder at that idx."""
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            ph.text_frame.text = text
            return ph
    return None


def add_bullets(slide, idx, lines, level_prefix="- "):
    """Fill a body/content placeholder with one paragraph per line in `lines`.
    Lines starting with a tab or 2 spaces become a level-1 (indented) bullet —
    mirrors how you'd type a nested list in Markdown."""
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            tf = ph.text_frame
            tf.clear()
            for i, line in enumerate(lines):
                level = 1 if line.startswith(("\t", "  ")) else 0
                text = line.strip().lstrip("-").strip()
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.text = text
                p.level = level
            return ph
    return None


def delete_slide(prs, index):
    """Remove a slide by index. python-pptx has no built-in delete — this drops
    it from both the slide-id list and the underlying XML relationship, which
    is what real deletion requires (removing just one leaves a dangling ref)."""
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    rId = slides[index].rId
    prs.part.drop_rel(rId)
    xml_slides.remove(slides[index])


def duplicate_slide(prs, index):
    """Duplicate slide at `index` (deep-copies its XML + images) and append it
    to the end. Use this to clone one of the two pre-built network-topology
    example slides (see references/template_map.md) as a starting point for a
    new diagram, instead of building one from blank."""
    from pptx.oxml.ns import qn

    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)

    # add_slide() pre-populates dest with empty placeholder shapes inherited
    # from the layout — drop them, we're about to copy the source's actual
    # (filled-in) shapes over instead.
    for shape in list(dest.shapes):
        shape._element.getparent().remove(shape._element)

    # Relationship rIds are only unique within one slide part, so copying XML
    # that still says r:embed="rId3" only works if the destination slide's
    # rId3 points at the same target — get_or_add() may hand back a different
    # rId, so track old->new and rewrite every r:* reference in the copied XML.
    rid_map = {}
    for rId, rel in source.part.rels.items():
        if "notesSlide" in rel.reltype:
            continue
        if rel.is_external:
            rid_map[rId] = dest.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            rid_map[rId] = dest.part.rels.get_or_add(rel.reltype, rel.target_part)

    r_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    for shape in source.shapes:
        new_el = copy.deepcopy(shape._element)
        for el in new_el.iter():
            for attr, old_rid in list(el.attrib.items()):
                if attr.startswith(f"{{{r_ns}}}") and old_rid in rid_map:
                    el.attrib[attr] = rid_map[old_rid]
        dest.shapes._spTree.append(new_el)

    return dest


def copy_shape(source_slide, dest_slide, shape_id=None, shape_name=None, offset_emu=(0, 0)):
    """Copy ONE shape (an icon from the icon-library slide, e.g.) from
    source_slide to dest_slide, optionally nudging it by (dx, dy) in EMU so
    copies don't stack exactly on top of each other. Match by shape_id (exact,
    from list_slides()) or shape_name (first match) — id is more reliable
    since many icons share the generic name 'Shape NNN'.
    """
    target = None
    for shp in source_slide.shapes:
        if shape_id is not None and shp.shape_id == shape_id:
            target = shp
            break
        if shape_name is not None and shp.name == shape_name:
            target = shp
            break
    if target is None:
        raise KeyError(f"no shape matched shape_id={shape_id} shape_name={shape_name}")

    new_el = copy.deepcopy(target._element)
    dest_slide.shapes._spTree.append(new_el)
    new_shape = dest_slide.shapes[-1]
    if offset_emu != (0, 0) and hasattr(new_shape, "left") and new_shape.left is not None:
        new_shape.left = Emu(new_shape.left + offset_emu[0])
        new_shape.top = Emu(new_shape.top + offset_emu[1])
    return new_shape
